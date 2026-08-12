from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from ozon_app.costs import CostCatalogError, build_cost_changes, export_cost_catalog, read_cost_catalog
from ozon_app.database import Database
from ozon_app.models import Product


class CostCatalogTests(unittest.TestCase):
    def test_export_read_preview_and_atomic_history(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_products = [
                Product("A-001", "Товар A", material_cost=80, labor_cost=20, active=True),
                Product("B-002", "Товар B", material_cost=45, labor_cost=5, active=False),
            ]
            path = export_cost_catalog(source_products, root / "costs.xlsx")
            imported = read_cost_catalog(path)

            self.assertEqual([product.article for product in imported], ["A-001", "B-002"])
            self.assertEqual(imported[0].total_cost, 100)
            self.assertEqual(imported[0].labor_cost, 20)
            self.assertFalse(imported[1].active)

            existing = {"A-001": Product("A-001", "Товар A", material_cost=70, labor_cost=20)}
            changes = build_cost_changes(imported, existing)
            self.assertEqual([change.status for change in changes], ["Изменение", "Новая позиция"])

            database = Database(root / "app.sqlite3")
            changed = database.save_products(imported, source="Тестовый импорт")
            self.assertEqual(changed, 2)
            history = database.list_product_cost_history()
            self.assertEqual(len(history), 2)
            self.assertTrue(all(row["change_source"] == "Тестовый импорт" for row in history))

    def test_rejects_labor_above_total_and_duplicate_article(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = export_cost_catalog([], Path(directory) / "invalid.xlsx")
            workbook = load_workbook(path)
            ws = workbook["Себестоимость"]
            ws["A5"], ws["B5"], ws["C5"], ws["D5"] = "A", "Товар", 100, 120
            ws["A6"], ws["B6"], ws["C6"], ws["D6"] = "A", "Дубль", 100, 20
            workbook.save(path)
            workbook.close()

            with self.assertRaises(CostCatalogError) as context:
                read_cost_catalog(path)
            self.assertIn("трудозатраты", str(context.exception))
            self.assertIn("уже указан", str(context.exception))


if __name__ == "__main__":
    unittest.main()
