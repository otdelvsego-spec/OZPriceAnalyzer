from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from ozon_app.database import Database
from ozon_app.exporter import export_calculation, export_run
from ozon_app.models import ParsedSource, Product, ProductResult, RunCalculation
from ozon_app.report_totals import report_total_value


class DatabaseExporterTests(unittest.TestCase):
    def test_history_snapshot_and_excel_breakdown(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            database = Database(root / "app.sqlite3")
            source_path = root / "source.xlsx"
            source = ParsedSource(
                path=source_path,
                file_hash="hash-1",
                report_type="ACCRUAL",
                sheet_name="Начисления",
                header_row=1,
            )
            calculation = RunCalculation(
                run_id=None,
                period_start=None,
                period_end=None,
                tax_rate=0.04,
                products=[
                    ProductResult(
                        article="A-1",
                        name="Товар",
                        material_cost=70,
                        labor_cost=30,
                        category="Старая категория",
                        units=2,
                        revenue_no_points=500,
                        commission=-100,
                        financial_result=400,
                    )
                ],
                unallocated_total=-50,
                unallocated={"Подписка Premium": (1, -50)},
                accrual_stats={"Подписка Premium": (0, 1)},
                source_files=[source],
            )
            run_id = database.save_run(calculation, {str(source_path): source_path})

            loaded = database.load_calculation(run_id)
            self.assertEqual(loaded.products[0].material_cost, 70)
            self.assertEqual(loaded.products[0].labor_cost, 30)
            self.assertEqual(loaded.products[0].category, "Старая категория")
            self.assertEqual(loaded.unallocated_total, -50)
            self.assertAlmostEqual(report_total_value(loaded), 130.0)
            summary = database.list_runs()[0]
            self.assertAlmostEqual(summary.commission_share, 0.2)
            self.assertEqual(summary.logistics_share, 0.0)
            self.assertEqual(summary.points_share, 0.0)
            self.assertAlmostEqual(summary.net_margin, 0.26)
            self.assertAlmostEqual(summary.profitability, 0.9)

            database.save_product(
                Product("A-1", "Товар", material_cost=999, labor_cost=1, category="Новая категория")
            )
            refreshed = database.load_calculation(run_id)
            self.assertEqual(refreshed.products[0].category, "Новая категория")
            self.assertEqual(refreshed.products[0].material_cost, 70)

            destination = root / "result.xlsx"
            export_run(database, run_id, destination)
            workbook = load_workbook(destination, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["КонсОтчет", "Разбивка"])
                sheet = workbook["КонсОтчет"]
                self.assertEqual(sheet["H4"].value, -50)
                self.assertEqual(sheet["L3"].value, "Чистая прибыль товаров")
                self.assertEqual(
                    sheet["M3"].value,
                    "Итог отчета с учетом нераспределенных после налога",
                )
                self.assertEqual(sheet["M4"].value, "=L7+H4-J4")
                self.assertEqual(sheet["K6"].value, "Финрезультат Ozon на ед.")
                self.assertEqual(sheet["M6"].value, "Финрезультат Ozon до с/с и налога")
                self.assertEqual(sheet["AH6"].value, "Средняя комиссия, % от выручки")
                self.assertEqual(sheet["AI6"].value, "Логистика, % от выручки")
                self.assertEqual(sheet["AJ6"].value, "Баллы, % от выручки")
                self.assertEqual(sheet["AK6"].value, "Чистая прибыль, % от выручки")
                self.assertEqual(sheet["AH7"].value, "=IFERROR(-V7/R7,0)")
                self.assertEqual(sheet["AI7"].value, "=IFERROR(-(Y7+Z7)/R7,0)")
                self.assertEqual(sheet["AJ7"].value, "=IFERROR(U7/R7,0)")
                self.assertEqual(sheet["AK7"].value, "=IFERROR((L7+$H$4-$J$4)/R7,0)")
                self.assertEqual(sheet["AH8"].value, "=IFERROR(-V8/R8,0)")
                self.assertEqual(sheet["AI8"].value, "=IFERROR(-(Y8+Z8)/R8,0)")
                self.assertEqual(sheet["AJ8"].value, "=IFERROR(U8/R8,0)")
                self.assertEqual(sheet["AK8"].value, "=IFERROR(L8/R8,0)")
                self.assertEqual(sheet["AH8"].number_format, "0.00%")
                self.assertEqual(sheet["AW6"].value, "Прибыль до себестоимости")
                self.assertEqual(workbook["Разбивка"]["A5"].value, "Подписка Premium")
                self.assertEqual(workbook["Разбивка"]["C5"].value, -50)
                self.assertTrue(str(sheet["AO8"].value).startswith("=IF"))
            finally:
                workbook.close()

    def test_compensation_tax_is_visible_in_history_and_excel_totals(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            database = Database(root / "app.sqlite3")
            source_path = root / "source.xlsx"
            source = ParsedSource(
                path=source_path,
                file_hash="hash-compensation",
                report_type="ACCRUAL",
                sheet_name="Начисления",
                header_row=1,
            )
            calculation = RunCalculation(
                run_id=None,
                period_start=None,
                period_end=None,
                tax_rate=0.04,
                products=[
                    ProductResult(
                        article="A-1",
                        name="Товар",
                        material_cost=100,
                        labor_cost=0,
                        units=1,
                        revenue_no_points=500,
                        financial_result=300,
                    )
                ],
                unallocated_total=120.78,
                unallocated={"Брак по вине Ozon на складе": (1, 120.78)},
                accrual_stats={"Брак по вине Ozon на складе": (0, 1)},
                source_files=[source],
            )
            run_id = database.save_run(calculation, {str(source_path): source_path})

            loaded = database.load_calculation(run_id)
            self.assertAlmostEqual(loaded.unallocated_compensation_income, 120.78)
            self.assertAlmostEqual(loaded.compensation_tax, 4.8312)
            self.assertAlmostEqual(report_total_value(loaded), 295.9488)
            self.assertAlmostEqual(database.list_runs()[0].net_margin, 295.9488 / 500)

            destination = root / "compensation.xlsx"
            export_run(database, run_id, destination)
            workbook = load_workbook(destination, data_only=False)
            try:
                sheet = workbook["КонсОтчет"]
                self.assertEqual(sheet["H4"].value, 120.78)
                self.assertEqual(sheet["I4"].value, 120.78)
                self.assertAlmostEqual(sheet["J4"].value, 4.8312)
                self.assertEqual(sheet["O7"].value, "=SUM(O8:O8)+$J$4")
                self.assertEqual(sheet["P7"].value, "=SUM(P8:P8)+$I$4")
                breakdown = workbook["Разбивка"]
                labels = [breakdown.cell(row, 1).value for row in range(1, breakdown.max_row + 1)]
                tax_row = labels.index("Налог с компенсаций") + 1
                self.assertAlmostEqual(breakdown.cell(tax_row, 3).value, 4.8312)
            finally:
                workbook.close()

    def test_aggregate_export_uses_exact_cost_and_tax_totals(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result = ProductResult(
                article="A-1",
                name="Товар",
                material_cost=72,
                labor_cost=28,
                units=5,
                revenue_no_points=1400,
                financial_result=900,
                material_sold_override=360,
                labor_sold_override=140,
                tax_override=74,
            )
            calculation = RunCalculation(
                run_id=None,
                period_start=None,
                period_end=None,
                tax_rate=74 / 1400,
                products=[result],
                unallocated_total=0,
                unallocated={},
                accrual_stats={},
            )
            destination = Path(directory) / "aggregate.xlsx"

            export_calculation(calculation, destination)

            workbook = load_workbook(destination, data_only=False)
            try:
                sheet = workbook["КонсОтчет"]
                self.assertEqual(sheet["F8"].value, 360)
                self.assertEqual(sheet["G8"].value, 140)
                self.assertEqual(sheet["H8"].value, "=F8+G8")
                self.assertEqual(sheet["O8"].value, 74)
                self.assertAlmostEqual(sheet["P4"].value, 74 / 1400)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
