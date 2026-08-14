from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


_TEMPLATE_NAMES = {"cost_template.xlsx", "report_template.xlsx"}


def ensure_clean_template(name: str, destination: str | Path) -> Path:
    """Create a data-free application template when it is requested."""
    if name not in _TEMPLATE_NAMES:
        raise ValueError(f"Unknown generated template: {name}")
    output = Path(destination).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    if name == "cost_template.xlsx":
        _build_cost_template(output)
    else:
        _build_report_template(output)
    return output


def _build_cost_template(output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Себестоимость"
    ws["A1"] = "OZ Price Analyzer — справочник себестоимости"
    ws.merge_cells("A1:H1")
    ws["A2"] = "Заполните товары начиная со строки 5. Шаблон не содержит рабочих данных."
    ws.merge_cells("A2:H2")

    headers = [
        "Артикул", "Наименование", "Категория", "Полная себестоимость, руб.",
        "Трудозатраты, руб.", "Материал, руб.", "Активен", "Комментарий",
    ]
    blue = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column, title in enumerate(headers, start=1):
        cell = ws.cell(4, column, title)
        cell.font = Font(bold=True)
        cell.fill = blue
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in range(5, 205):
        for column in range(1, 9):
            ws.cell(row, column).border = border
        ws.cell(row, 6).value = f'=IF(OR(D{row}="",E{row}=""),"",D{row}-E{row})'
        for column in (4, 5, 6):
            ws.cell(row, column).number_format = "#,##0.00"

    table = Table(displayName="CostCatalog", ref="A4:H204")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    active_validation = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=True)
    active_validation.add("G5:G204")
    ws.add_data_validation(active_validation)
    for column, width in {
        "A": 18, "B": 34, "C": 24, "D": 24,
        "E": 22, "F": 20, "G": 12, "H": 32,
    }.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A5"
    wb.save(output)
    wb.close()


def _build_report_template(output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "КонсОтчет"
    ws["A1"] = "OZ Price Analyzer — итоговый отчет"
    ws["A1"].font = Font(size=16, bold=True)
    ws["B5"] = "Период: будет заполнен при экспорте"
    ws["H3"] = "Нераспределенные доходы / расходы"
    ws["P3"] = "Налоговая ставка"

    actual_headers = [
        "Артикул", "Наименование", "Итого с/с", "Материал", "Трудозатраты",
        "Материал проданного", "Трудозатраты проданного", "С/с проданного",
        "Доходность", "Чистая прибыль на ед.", "Прибыль от продаж на ед.",
        "Чистая прибыль всего", "Прибыль от продаж всего", "Средняя цена",
        "Налог", "Налогооблагаемый доход", "Продажи", "Выручка с баллами",
        "Выручка без баллов", "Программы партнеров", "Баллы", "Комиссия Ozon",
        "Обработка отправления", "Доставка до ПВЗ", "Логистика",
        "Обратная логистика", "Возвраты/отмены", "Эквайринг", "Звездные товары",
        "Упаковка и материалы", "Компенсации Ozon", "Прочие начисления",
        "Финрезультат Ozon",
    ]
    for column, title in enumerate(actual_headers, start=1):
        ws.cell(6, column, title)

    scenario_headers = {
        38: "Текущая средняя цена", 39: "Изменение цены", 40: "Плановая цена",
        41: "Плановая доходность", 42: "Затраты Ozon без комиссии",
        43: "Плановая выручка", 44: "Средняя комиссия", 45: "Плановая комиссия",
        46: "Плановые баллы", 47: "Налоговая база", 48: "Налог",
        49: "Прибыль от продаж", 50: "Прибыль/ед. до с/с", 51: "Чистая прибыль/ед.",
    }
    for column, title in scenario_headers.items():
        ws.cell(6, column, title)

    blue = PatternFill("solid", fgColor="D9EAF7")
    green = PatternFill("solid", fgColor="E2F0D9")
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column in range(1, 52):
        for row in (6, 7, 8):
            cell = ws.cell(row, column)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=row == 6)
        ws.cell(6, column).font = Font(bold=True)
        ws.cell(6, column).fill = blue
        ws.cell(7, column).fill = green
    for row in range(9, 29):
        for column in range(1, 52):
            ws.cell(row, column)._style = ws.cell(8, column)._style

    for row in range(7, 29):
        for column in list(range(3, 17)) + list(range(18, 34)) + list(range(38, 52)):
            ws.cell(row, column).number_format = "#,##0.00"
        for column in (9, 39, 41, 44):
            ws.cell(row, column).number_format = "0.00%"

    for column in range(1, 52):
        ws.column_dimensions[get_column_letter(column)].width = 16
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 34
    ws.freeze_panes = "B1"
    wb.save(output)
    wb.close()
