from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

from .config import resource_path
from .database import Database
from .models import ProductResult, RunCalculation


RESULT_COLUMNS = {
    "units": "Q",
    "revenue_no_points": "S",
    "partner_programs": "T",
    "points": "U",
    "commission": "V",
    "processing": "W",
    "delivery": "X",
    "logistics": "Y",
    "reverse_logistics": "Z",
    "returns_cancels": "AA",
    "acquiring": "AB",
    "stars": "AC",
    "packaging": "AD",
    "compensation": "AE",
    "other": "AF",
}


def export_run(database: Database, run_id: int, destination: str | Path) -> Path:
    calculation = database.load_calculation(run_id)
    planned_prices = database.planned_prices(run_id)
    return export_calculation(calculation, destination, planned_prices)


def export_calculation(
    calculation: RunCalculation,
    destination: str | Path,
    planned_prices: dict[str, float] | None = None,
) -> Path:
    template = resource_path("report_template.xlsx")
    if not template.exists():
        raise FileNotFoundError("Не найден шаблон итогового отчета")
    workbook = load_workbook(template)
    if "Справочник начислений" in workbook.sheetnames:
        del workbook["Справочник начислений"]
    ws = workbook["КонсОтчет"]
    _fill_report_sheet(ws, calculation, planned_prices or {})
    _create_breakdown_sheet(workbook, calculation)
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    output = Path(destination).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    return output


def suggested_export_name(calculation: RunCalculation) -> str:
    if calculation.period_start and calculation.period_end:
        period = f"{calculation.period_start:%d.%m.%Y}-{calculation.period_end:%d.%m.%Y}"
    else:
        period = datetime.now().strftime("%d.%m.%Y")
    return f"Отчет_OZON_{period}.xlsx"


def _fill_report_sheet(ws, calculation: RunCalculation, planned_prices: dict[str, float]) -> None:
    template_last_row = max(ws.max_row, 28)
    output_last_row = 7 + len(calculation.products)
    clear_to = max(template_last_row, output_last_row)
    for row_number in range(8, clear_to + 1):
        if row_number != 8:
            _copy_row_style(ws, 8, row_number, 51)
        for column in range(1, 52):
            ws.cell(row_number, column).value = None

    for index, result in enumerate(calculation.products, start=8):
        _write_product_row(ws, index, result, calculation.tax_rate, planned_prices.get(result.article))

    last_row = max(output_last_row, 8)
    ws["H4"] = calculation.unallocated_total
    ws["I4"] = calculation.unallocated_compensation_income
    ws["J4"] = calculation.compensation_tax
    ws["P4"] = calculation.tax_rate
    ws["M4"] = "=L7+H4-J4"
    if calculation.period_start and calculation.period_end:
        ws["B5"] = f"Период: {calculation.period_start:%d.%m.%Y}-{calculation.period_end:%d.%m.%Y}"
    else:
        ws["B5"] = "Период не определен"
    for column in ("F", "G", "H", "L", "M", "O", "P"):
        ws[f"{column}7"] = f"=SUM({column}8:{column}{last_row})"
    ws["O7"] = f"=SUM(O8:O{last_row})+$J$4"
    ws["P7"] = f"=SUM(P8:P{last_row})+$I$4"
    for column_number in range(17, 34):
        letter = ws.cell(1, column_number).column_letter
        ws[f"{letter}7"] = f"=SUM({letter}8:{letter}{last_row})"
    ws["AH7"] = "=IFERROR(-V7/R7,0)"
    ws["AI7"] = "=IFERROR(-(Y7+Z7)/R7,0)"
    ws["AJ7"] = "=IFERROR(U7/R7,0)"
    ws["AK7"] = "=IFERROR((L7+$H$4-$J$4)/R7,0)"
    ws["AL7"] = "=IFERROR(R7/Q7,\"\")"
    ws["AM7"] = "=IF(OR(AL7=\"\",AN7=\"\"),\"\",IFERROR(AN7/AL7-1,0))"
    ws["AN7"] = "=IFERROR(AQ7/Q7,\"\")"
    ws["AO7"] = "=IFERROR((AW7-H7)/H7,0)"
    for column in ("AP", "AQ", "AS", "AT", "AU", "AV", "AW"):
        ws[f"{column}7"] = f"=SUM({column}8:{column}{last_row})"
    ws["AR7"] = "=IFERROR(-V7/R7,0)"
    ws["AX7"] = "=IFERROR(AW7/Q7,0)"
    ws["AY7"] = "=IFERROR((AW7-H7)/Q7,0)"
    ws.freeze_panes = "B1"


def _write_product_row(
    ws,
    row: int,
    result: ProductResult,
    tax_rate: float,
    planned_price: float | None,
) -> None:
    ws[f"A{row}"] = result.article
    ws[f"B{row}"] = result.name
    ws[f"D{row}"] = result.material_cost
    ws[f"E{row}"] = result.labor_cost
    ws[f"C{row}"] = f"=D{row}+E{row}"
    ws[f"F{row}"] = (
        result.material_sold
        if result.material_sold_override is not None
        else f"=D{row}*Q{row}"
    )
    ws[f"G{row}"] = (
        result.labor_sold
        if result.labor_sold_override is not None
        else f"=E{row}*Q{row}"
    )
    ws[f"H{row}"] = f"=F{row}+G{row}"
    ws[f"I{row}"] = f"=IFERROR(J{row}/C{row},0)"
    ws[f"J{row}"] = f"=IFERROR(L{row}/Q{row},0)"
    ws[f"K{row}"] = f"=IFERROR(M{row}/Q{row},0)"
    ws[f"L{row}"] = f"=M{row}-H{row}-O{row}"
    ws[f"M{row}"] = f"=AG{row}"
    ws[f"N{row}"] = f"=IFERROR(R{row}/Q{row},\"\")"
    ws[f"O{row}"] = (
        result.tax(tax_rate)
        if result.tax_override is not None
        else f"=P{row}*$P$4"
    )
    ws[f"P{row}"] = f"=S{row}+T{row}"
    ws[f"Q{row}"] = result.units
    ws[f"R{row}"] = f"=SUM(S{row}:U{row})"
    for field, column in RESULT_COLUMNS.items():
        ws[f"{column}{row}"] = getattr(result, field)
    ws[f"AG{row}"] = f"=SUM(S{row}:AF{row})"
    ws[f"AH{row}"] = f"=IFERROR(-V{row}/R{row},0)"
    ws[f"AI{row}"] = f"=IFERROR(-(Y{row}+Z{row})/R{row},0)"
    ws[f"AJ{row}"] = f"=IFERROR(U{row}/R{row},0)"
    ws[f"AK{row}"] = f"=IFERROR(L{row}/R{row},0)"

    ws[f"AL{row}"] = f"=IFERROR(R{row}/Q{row},\"\")"
    ws[f"AM{row}"] = f"=IF(OR(AL{row}=\"\",AN{row}=\"\"),\"\",IFERROR(AN{row}/AL{row}-1,0))"
    ws[f"AN{row}"] = planned_price if planned_price is not None else result.average_price()
    ws[f"AO{row}"] = f"=IF(OR(Q{row}=0,AN{row}=\"\"),\"\",IFERROR(AY{row}/C{row},0))"
    ws[f"AP{row}"] = f"=IF(Q{row}=0,\"\",R{row}-M{row}+V{row})"
    ws[f"AQ{row}"] = f"=IF(OR(AN{row}=\"\",Q{row}=0),\"\",AN{row}*Q{row})"
    ws[f"AR{row}"] = f"=IF(OR(R{row}=0,Q{row}=0),\"\",IFERROR(-V{row}/R{row},0))"
    ws[f"AS{row}"] = f"=IF(AQ{row}=\"\",\"\",-AQ{row}*AR{row})"
    ws[f"AT{row}"] = f"=IF(OR(AQ{row}=\"\",R{row}=0),\"\",IFERROR(U{row}/R{row}*AQ{row},0))"
    ws[f"AU{row}"] = f"=IF(AQ{row}=\"\",\"\",AQ{row}-AT{row})"
    ws[f"AV{row}"] = f"=IF(AU{row}=\"\",\"\",AU{row}*$P$4)"
    ws[f"AW{row}"] = f"=IF(AQ{row}=\"\",\"\",AQ{row}-AP{row}+AS{row}-AV{row})"
    ws[f"AX{row}"] = f"=IF(OR(AW{row}=\"\",Q{row}=0),\"\",AW{row}/Q{row})"
    ws[f"AY{row}"] = f"=IF(AX{row}=\"\",\"\",AX{row}-C{row})"


def _copy_row_style(ws, source_row: int, target_row: int, last_column: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for column in range(1, last_column + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def _create_breakdown_sheet(workbook, calculation: RunCalculation) -> None:
    if "Разбивка" in workbook.sheetnames:
        del workbook["Разбивка"]
    ws = workbook.create_sheet("Разбивка")
    ws.merge_cells("A1:C1")
    ws["A1"] = "Нераспределенные доходы / расходы по типам начисления"
    ws["A2"] = "Период"
    if calculation.period_start and calculation.period_end:
        ws["B2"] = f"{calculation.period_start:%d.%m.%Y}-{calculation.period_end:%d.%m.%Y}"
    ws.append([])
    ws.append(["Тип начисления", "Количество строк", "Сумма, руб."])
    for accrual_type, (row_count, amount) in calculation.unallocated.items():
        ws.append([accrual_type, row_count, amount])
    total_row = ws.max_row + 1
    ws.cell(total_row, 1, "Итого")
    if total_row > 5:
        ws.cell(total_row, 2, f"=SUM(B5:B{total_row - 1})")
        ws.cell(total_row, 3, f"=SUM(C5:C{total_row - 1})")
    else:
        ws.cell(total_row, 2, 0)
        ws.cell(total_row, 3, 0)
    ws.cell(total_row + 2, 1, "Контроль: сумма нераспределенных")
    ws.cell(total_row + 2, 3, "='КонсОтчет'!H4")
    ws.cell(total_row + 3, 1, "Отклонение")
    ws.cell(total_row + 3, 3, f"=C{total_row}-C{total_row + 2}")
    ws.cell(total_row + 5, 1, "Компенсации в налогооблагаемой базе")
    ws.cell(total_row + 5, 3, calculation.unallocated_compensation_income)
    ws.cell(total_row + 6, 1, "Налог с компенсаций")
    ws.cell(total_row + 6, 3, calculation.compensation_tax)
    ws.cell(total_row + 7, 1, "Нераспределенные после налога с компенсаций")
    ws.cell(total_row + 7, 3, calculation.unallocated_total - calculation.compensation_tax)

    blue = PatternFill("solid", fgColor="D9E1F2")
    green = PatternFill("solid", fgColor="E2EFDA")
    thin = Side(style="thin", color="7F7F7F")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = blue
        cell.font = Font(name="Calibri", size=14, bold=True)
        cell.alignment = Alignment(horizontal="center")
    for cell in ws[4]:
        cell.fill = blue
        cell.font = Font(bold=True)
        cell.border = border
    for row in ws.iter_rows(min_row=5, max_row=total_row, min_col=1, max_col=3):
        for cell in row:
            cell.border = border
    for cell in ws[total_row]:
        cell.fill = green
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    for row in range(5, total_row + 8):
        ws.cell(row, 3).number_format = "#,##0.00"
    ws.auto_filter.ref = f"A4:C{total_row}"
    ws.freeze_panes = "A5"
