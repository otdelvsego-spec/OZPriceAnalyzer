from __future__ import annotations

import hashlib
import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import AccrualRow, ParsedSource, RealizationRow


REPORT_ACCRUAL = "ACCRUAL"
REPORT_REALIZATION = "REALIZATION"


class ReportFormatError(ValueError):
    """Raised when an XLSX file is not a supported or valid Ozon report."""


def normalize_text(value: object) -> str:
    text = "" if value is None else str(value)
    return " ".join(text.replace("\r", " ").replace("\n", " ").strip().split()).casefold()


def display_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_sku(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return str(int(float(value))) if float(value).is_integer() else format(float(value), "f").rstrip("0").rstrip(".")
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def as_float(value: object, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return default
    if isinstance(value, (int, float)):
        result = float(value)
        return result if math.isfinite(result) else default
    text = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".").strip()
    try:
        result = float(text)
    except ValueError:
        return default
    return result if math.isfinite(result) else default


def is_numeric(value: object) -> bool:
    if value is None or value == "" or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return math.isfinite(float(value))
    text = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".").strip()
    try:
        return math.isfinite(float(text))
    except ValueError:
        return False


def as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _row_caption_map(ws, row_number: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in ws[row_number]:
        key = normalize_text(cell.value)
        if key and key not in result:
            result[key] = cell.column
    return result


def _sheet_max_row(ws) -> int:
    """Return a reliable row count even when Ozon omitted sheet dimensions."""
    if ws.max_row is None:
        calculate_dimension = getattr(ws, "calculate_dimension", None)
        if calculate_dimension is not None:
            calculate_dimension(force=True)
    return int(ws.max_row or 0)


def _sheet_max_column(ws) -> int:
    if ws.max_column is None:
        calculate_dimension = getattr(ws, "calculate_dimension", None)
        if calculate_dimension is not None:
            calculate_dimension(force=True)
    return int(ws.max_column or 0)


def _find_exact(columns: dict[str, int], caption: str) -> int:
    return columns.get(normalize_text(caption), 0)


def _find_in_rows(ws, first_row: int, last_row: int, caption: str) -> int:
    for row_number in range(first_row, last_row + 1):
        column = _find_exact(_row_caption_map(ws, row_number), caption)
        if column:
            return column
    return 0


def _detect_sheet(workbook) -> tuple[str, object, int]:
    for ws in workbook.worksheets:
        max_row = _sheet_max_row(ws)
        for row_number in range(1, min(max_row, 10) + 1):
            columns = _row_caption_map(ws, row_number)
            if all(_find_exact(columns, name) for name in ("Тип начисления", "Артикул", "Сумма итого, руб.")):
                return REPORT_ACCRUAL, ws, row_number
    for ws in workbook.worksheets:
        max_row = _sheet_max_row(ws)
        for row_number in range(1, min(max_row, 20) + 1):
            columns = _row_caption_map(ws, row_number)
            if all(_find_exact(columns, name) for name in ("Код товара продавца", "Код товара OZON", "Номер отправления")):
                return REPORT_REALIZATION, ws, row_number
    raise ReportFormatError(
        "Тип файла не определен. Выберите отчет по начислениям или RealizationReportCIS. "
        "Отдельный CompensationReport не используется."
    )


def parse_report(path: str | Path) -> ParsedSource:
    source_path = Path(path).expanduser().resolve()
    if source_path.suffix.casefold() != ".xlsx":
        raise ReportFormatError(f"Поддерживаются только файлы XLSX: {source_path.name}")
    try:
        # Ozon exports often omit worksheet dimensions. Normal mode is both
        # reliable and much faster here than random cell access in read-only mode.
        workbook = load_workbook(source_path, read_only=False, data_only=True)
    except Exception as exc:
        raise ReportFormatError(f"Не удалось прочитать {source_path.name}: {exc}") from exc
    try:
        report_type, ws, header_row = _detect_sheet(workbook)
        parsed = ParsedSource(
            path=source_path,
            file_hash=sha256_file(source_path),
            report_type=report_type,
            sheet_name=ws.title,
            header_row=header_row,
        )
        if report_type == REPORT_ACCRUAL:
            _parse_accrual(ws, header_row, parsed)
        else:
            _parse_realization(ws, header_row, parsed)
        return parsed
    finally:
        workbook.close()


def _parse_accrual(ws, header_row: int, parsed: ParsedSource) -> None:
    columns = _row_caption_map(ws, header_row)
    required = {
        "id": "ID начисления",
        "date": "Дата начисления",
        "type": "Тип начисления",
        "article": "Артикул",
        "quantity": "Количество",
        "price": "Цена продавца",
        "amount": "Сумма итого, руб.",
    }
    positions = {key: _find_exact(columns, caption) for key, caption in required.items()}
    missing = [required[key] for key, column in positions.items() if not column]
    if missing:
        raise ReportFormatError(
            f"В отчете {parsed.path.name} изменились обязательные заголовки: {', '.join(missing)}"
        )
    sku_column = _find_exact(columns, "SKU")
    name_column = _find_exact(columns, "Название товара")
    dates: list[date] = []
    for row_number in range(header_row + 1, _sheet_max_row(ws) + 1):
        accrual_type = display_text(ws.cell(row_number, positions["type"]).value)
        if not accrual_type:
            continue
        accrual_date = as_date(ws.cell(row_number, positions["date"]).value)
        if accrual_date:
            dates.append(accrual_date)
        parsed.accrual_rows.append(
            AccrualRow(
                source_name=parsed.path.name,
                sheet_name=ws.title,
                row_number=row_number,
                accrual_id=display_text(ws.cell(row_number, positions["id"]).value),
                accrual_date=accrual_date,
                accrual_type=accrual_type,
                article=display_text(ws.cell(row_number, positions["article"]).value),
                sku=normalize_sku(ws.cell(row_number, sku_column).value) if sku_column else "",
                product_name=display_text(ws.cell(row_number, name_column).value) if name_column else "",
                quantity=as_float(ws.cell(row_number, positions["quantity"]).value),
                seller_price=as_float(ws.cell(row_number, positions["price"]).value),
                amount=as_float(ws.cell(row_number, positions["amount"]).value),
            )
        )
    if not parsed.accrual_rows:
        raise ReportFormatError(f"В отчете по начислениям нет строк данных: {parsed.path.name}")
    parsed.period_start = min(dates) if dates else None
    parsed.period_end = max(dates) if dates else None


def _parse_realization(ws, header_row: int, parsed: ParsedSource) -> None:
    columns = _row_caption_map(ws, header_row)
    article_column = _find_exact(columns, "Код товара продавца")
    sku_column = _find_exact(columns, "Код товара OZON")
    name_column = _find_exact(columns, "Товар")
    shipment_column = _find_exact(columns, "Номер отправления")
    max_row = _sheet_max_row(ws)
    unit_price_column = _find_in_rows(ws, header_row, min(header_row + 2, max_row), "Цена реализации с НДС, руб.")
    quantity_column = _find_in_rows(ws, header_row, min(header_row + 2, max_row), "Кол-во")
    total_column = _find_in_rows(ws, header_row, min(header_row + 2, max_row), "Итого к начислению, руб.")
    required_positions = [article_column, sku_column, shipment_column, unit_price_column, quantity_column, total_column]
    if not all(required_positions):
        raise ReportFormatError(f"В отчете о выкупленных товарах изменились заголовки: {parsed.path.name}")

    parsed.period_start, parsed.period_end = _realization_period(ws, header_row)

    detail_total = 0.0
    for row_number in range(header_row + 1, max_row + 1):
        raw_article = display_text(ws.cell(row_number, article_column).value)
        if not _is_realization_detail(ws, row_number, raw_article, quantity_column, total_column):
            continue
        quantity = as_float(ws.cell(row_number, quantity_column).value)
        unit_price = as_float(ws.cell(row_number, unit_price_column).value)
        amount = as_float(ws.cell(row_number, total_column).value)
        if quantity <= 0 or abs(amount - unit_price * quantity) > 0.05:
            raise ReportFormatError(
                f"Поврежден или изменен отчет {parsed.path.name}: строка {row_number}, "
                "итог не равен цене реализации, умноженной на количество."
            )
        detail_total += amount
        parsed.realization_rows.append(
            RealizationRow(
                source_name=parsed.path.name,
                sheet_name=ws.title,
                row_number=row_number,
                raw_article=raw_article,
                sku=normalize_sku(ws.cell(row_number, sku_column).value),
                product_name=display_text(ws.cell(row_number, name_column).value) if name_column else "",
                shipment=display_text(ws.cell(row_number, shipment_column).value),
                unit_price=unit_price,
                quantity=quantity,
                amount=amount,
            )
        )
    if not parsed.realization_rows:
        raise ReportFormatError(f"В отчете нет строк выкупленных товаров: {parsed.path.name}")

    stated_total = _find_stated_total(ws, total_column)
    if stated_total != 0 and abs(stated_total - detail_total) > 0.05:
        raise ReportFormatError(
            f"Поврежден отчет {parsed.path.name}: сумма строк {detail_total:,.2f} не совпадает "
            f"с итогом отчета {stated_total:,.2f}."
        )


def _realization_period(ws, header_row: int) -> tuple[date | None, date | None]:
    period_pattern = re.compile(
        r"за\s+период\s+с\s+(\d{1,2}\.\d{1,2}\.\d{4})\s+по\s+(\d{1,2}\.\d{1,2}\.\d{4})",
        flags=re.IGNORECASE,
    )
    for row_number in range(1, header_row):
        for column_number in range(1, _sheet_max_column(ws) + 1):
            value = display_text(ws.cell(row_number, column_number).value)
            match = period_pattern.search(value)
            if not match:
                continue
            start = as_date(match.group(1))
            end = as_date(match.group(2))
            if start is not None and end is not None and start <= end:
                return start, end
    return None, None


def _is_realization_detail(ws, row_number: int, raw_article: str, quantity_column: int, total_column: int) -> bool:
    second_column_value = ws.cell(row_number, 2).value
    return (
        is_numeric(second_column_value)
        and bool(raw_article)
        and not is_numeric(raw_article)
        and is_numeric(ws.cell(row_number, quantity_column).value)
        and is_numeric(ws.cell(row_number, total_column).value)
    )


def _find_stated_total(ws, total_column: int) -> float:
    for row_number in range(1, _sheet_max_row(ws) + 1):
        label = normalize_text(ws.cell(row_number, 2).value)
        value = ws.cell(row_number, total_column).value
        if ("итого" in label or "всего" in label) and is_numeric(value):
            return as_float(value)
    return 0.0


def workbook_sheet_names(path: str | Path) -> list[str]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def preview_sheet(
    path: str | Path,
    sheet_name: str,
    max_rows: int = 500,
    max_columns: int = 40,
) -> tuple[list[str], list[list[str]]]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        ws = workbook[sheet_name]
        row_count = _sheet_max_row(ws)
        column_count = min(_sheet_max_column(ws), max_columns)
        headers = ["Строка"] + [get_column_letter(index) for index in range(1, column_count + 1)]
        rows: list[list[str]] = []
        for row_index, values in enumerate(
            ws.iter_rows(min_row=1, max_row=min(row_count, max_rows), max_col=column_count, values_only=True),
            start=1,
        ):
            rows.append([str(row_index)] + [_format_preview_value(value) for value in values])
        return headers, rows
    finally:
        workbook.close()


def _format_preview_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.6f}".rstrip("0").rstrip(".")
    return str(value)


def all_rows(sources: Iterable[ParsedSource]) -> tuple[list[AccrualRow], list[RealizationRow]]:
    accrual: list[AccrualRow] = []
    realization: list[RealizationRow] = []
    for source in sources:
        accrual.extend(source.accrual_rows)
        realization.extend(source.realization_rows)
    return accrual, realization
