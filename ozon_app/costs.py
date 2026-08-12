from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .config import resource_path
from .excel_reader import as_float, display_text, is_numeric, normalize_text
from .models import Product


class CostCatalogError(ValueError):
    """Raised when a cost catalog cannot be imported safely."""


@dataclass(slots=True)
class CostChange:
    product: Product
    previous: Product | None
    status: str

    @property
    def changed(self) -> bool:
        return self.status != "Без изменений"

    @property
    def previous_total(self) -> float | None:
        return self.previous.total_cost if self.previous else None

    @property
    def total_change(self) -> float | None:
        if self.previous is None:
            return None
        return self.product.total_cost - self.previous.total_cost


@dataclass(slots=True)
class CostEditorEntry:
    article: object
    name: object
    total_cost: object
    labor_cost: object
    active: bool
    row_number: int


def build_products_from_editor_entries(entries: Iterable[CostEditorEntry]) -> list[Product]:
    products: list[Product] = []
    seen: dict[str, int] = {}
    errors: list[str] = []
    for entry in entries:
        article = display_text(entry.article)
        name = display_text(entry.name)
        total_value = entry.total_cost
        labor_value = entry.labor_cost
        if not article and not name and total_value in (None, "") and labor_value in (None, ""):
            continue
        location = f"строка {entry.row_number}"
        if not article:
            errors.append(f"{location}: не указан артикул")
            continue
        if article in seen:
            errors.append(f"{location}: артикул {article} уже указан в строке {seen[article]}")
            continue
        seen[article] = entry.row_number
        if not is_numeric(total_value):
            errors.append(f"{location}: не указана полная себестоимость для {article}")
            continue
        if labor_value not in (None, "") and not is_numeric(labor_value):
            errors.append(f"{location}: неверно указаны трудозатраты для {article}")
            continue
        total = as_float(total_value)
        labor = as_float(labor_value) if labor_value not in (None, "") else 0.0
        if total < 0 or labor < 0 or labor > total:
            errors.append(
                f"{location}: трудозатраты должны быть от 0 до полной себестоимости для {article}"
            )
            continue
        products.append(
            Product(
                article=article,
                name=name or article,
                material_cost=total - labor,
                labor_cost=labor,
                active=entry.active,
            )
        )
    if errors:
        detail = "\n".join(f"• {message}" for message in errors[:12])
        suffix = f"\n…и еще {len(errors) - 12}" if len(errors) > 12 else ""
        raise CostCatalogError(f"Справочник содержит ошибки:\n{detail}{suffix}")
    if not products:
        raise CostCatalogError("В справочнике нет заполненных товаров")
    return products


def read_cost_catalog(path: str | Path) -> list[Product]:
    source = Path(path).expanduser().resolve()
    if source.suffix.casefold() != ".xlsx":
        raise CostCatalogError("Справочник себестоимости должен быть в формате XLSX")
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise CostCatalogError(f"Не удалось прочитать справочник: {exc}") from exc
    try:
        ws, header_row, columns = _find_catalog_sheet(workbook)
        positions = {
            "article": columns[normalize_text("Артикул")],
            "name": columns[normalize_text("Наименование")],
            "total": columns[normalize_text("Полная себестоимость, руб.")],
            "labor": columns[normalize_text("Трудозатраты, руб.")],
        }
        active_column = columns.get(normalize_text("Активен"), 0)
        entries: list[CostEditorEntry] = []
        for row_number in range(header_row + 1, int(ws.max_row or header_row) + 1):
            article = display_text(ws.cell(row_number, positions["article"]).value)
            name = display_text(ws.cell(row_number, positions["name"]).value)
            total_value = ws.cell(row_number, positions["total"]).value
            labor_value = ws.cell(row_number, positions["labor"]).value
            if not article and not name and total_value in (None, "") and labor_value in (None, ""):
                continue
            entries.append(
                CostEditorEntry(
                    article=article,
                    name=name,
                    total_cost=total_value,
                    labor_cost=labor_value,
                    active=_active_value(ws.cell(row_number, active_column).value) if active_column else True,
                    row_number=row_number,
                )
            )
        return build_products_from_editor_entries(entries)
    finally:
        workbook.close()


def build_cost_changes(products: list[Product], existing: dict[str, Product]) -> list[CostChange]:
    changes: list[CostChange] = []
    for product in products:
        previous = existing.get(product.article)
        if previous is None:
            status = "Новая позиция"
        elif _same_product(previous, product):
            status = "Без изменений"
        else:
            status = "Изменение"
        changes.append(CostChange(product=product, previous=previous, status=status))
    return changes


def export_cost_catalog(products: list[Product], destination: str | Path) -> Path:
    template = resource_path("cost_template.xlsx")
    if not template.exists():
        raise FileNotFoundError("Не найден шаблон справочника себестоимости")
    workbook = load_workbook(template)
    ws = workbook["Себестоимость"]
    reserved_last_row = max(204, 4 + len(products))
    for row_number in range(5, reserved_last_row + 1):
        for column in (1, 2, 3, 4, 6, 7):
            ws.cell(row_number, column).value = None
        ws.cell(row_number, 5).value = f'=IF(OR(C{row_number}="",D{row_number}=""),"",C{row_number}-D{row_number})'
    for row_number, product in enumerate(products, start=5):
        ws.cell(row_number, 1).value = product.article
        ws.cell(row_number, 2).value = product.name
        ws.cell(row_number, 3).value = product.total_cost
        ws.cell(row_number, 4).value = product.labor_cost
        ws.cell(row_number, 6).value = "Да" if product.active else "Нет"
    ws["A3"] = f"Выгружено: {datetime.now():%d.%m.%Y %H:%M}"
    if ws.tables:
        next(iter(ws.tables.values())).ref = f"A4:G{reserved_last_row}"
    output = Path(destination).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    return output


def _find_catalog_sheet(workbook):
    required = {
        normalize_text("Артикул"),
        normalize_text("Наименование"),
        normalize_text("Полная себестоимость, руб."),
        normalize_text("Трудозатраты, руб."),
    }
    for ws in workbook.worksheets:
        for row_number in range(1, min(int(ws.max_row or 0), 15) + 1):
            columns = {
                normalize_text(cell.value): cell.column
                for cell in ws[row_number]
                if normalize_text(cell.value)
            }
            if required.issubset(columns):
                return ws, row_number, columns
    raise CostCatalogError(
        "Не найдены обязательные столбцы: Артикул, Наименование, "
        "Полная себестоимость, руб., Трудозатраты, руб."
    )


def _active_value(value: object) -> bool:
    return normalize_text(value) not in {"нет", "0", "false", "архив", "неактивен"}


def _same_product(first: Product, second: Product) -> bool:
    return (
        first.name.strip() == second.name.strip()
        and abs(first.material_cost - second.material_cost) < 0.005
        and abs(first.labor_cost - second.labor_cost) < 0.005
        and first.active == second.active
    )
